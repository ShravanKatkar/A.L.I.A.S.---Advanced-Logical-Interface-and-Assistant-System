import os
import re
import json
import time
from datetime import datetime
from pathlib import Path

# PDF Generation via ReportLab
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

# Word Generation via python-docx
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

class DocumentOps:
    """
    Automated professional document generation engine for ALIAS.
    Generates styled PDF and Microsoft Word (.docx) documents,
    saves them to the user's Desktop/Documents, and opens them automatically.
    """

    def __init__(self):
        # Preferred user output folders
        self.desktop_dir = os.path.join(os.path.expanduser("~"), "Desktop")
        self.documents_dir = os.path.join(os.path.expanduser("~"), "Documents")
        
        # Local upload backup directory for web serving / downloads
        self.local_docs_dir = Path(__file__).parent.parent / "uploads" / "documents"
        os.makedirs(self.local_docs_dir, exist_ok=True)

    def get_output_directory(self) -> str:
        """Returns the best directory to save user documents (Desktop if writable, else Documents)."""
        if os.path.exists(self.desktop_dir) and os.access(self.desktop_dir, os.W_OK):
            return self.desktop_dir
        if os.path.exists(self.documents_dir) and os.access(self.documents_dir, os.W_OK):
            return self.documents_dir
        return str(self.local_docs_dir)

    def detect_document_intent(self, text: str):
        """
        Detects if the user wants to generate a PDF or Word document.
        Returns a dict with doc_type and topic, or None.
        """
        if not text:
            return None

        clean = text.strip()
        lower = clean.lower()
        lower = re.sub(r'^(can you |could you |please |hey alias |alias |would you kindly )+', '', lower).strip()
        lower = re.sub(r'[.!?]+$', '', lower).strip()

        # Check for document format keywords
        is_pdf = bool(re.search(r'\b(pdf|pdf document|pdf file|pdf report)\b', lower))
        is_word = bool(re.search(r'\b(word|docx|word doc|word document|word file|ms word|microsoft word)\b', lower))

        if not (is_pdf or is_word):
            # Also check generic "create a document about X" -> default to PDF
            if re.search(r'\b(create|make|generate|build|write|draft|prepare)\s+(?:a|an)\s+(?:document|report|summary|handbook|guide|notes)\b', lower):
                is_pdf = True
            else:
                return None

        # Check for creation verbs
        create_pattern = r'\b(create|make|generate|build|write|draft|produce|prepare|save as)\b'
        if not re.search(create_pattern, lower):
            return None

        doc_type = "pdf" if is_pdf else "word"

        # Extract topic by stripping creation filler phrases
        topic = re.sub(
            r'^(?:create|make|generate|build|write|draft|produce|prepare)\s+(?:a|an|me a)?\s*(?:new)?\s*(?:pdf|word|docx|doc|ms word)?\s*(?:document|file|report|summary|handbook|notes|guide)?\s*(?:in\s+(?:pdf|word|docx)\s+format)?\s*(?:about|on|for|regarding|titled|explaining|covering)?\s*',
            '',
            lower
        ).strip()

        # Clean trailing filler
        topic = re.sub(r'\s+(in\s+pdf|in\s+word|as\s+pdf|as\s+word|in\s+docx)$', '', topic).strip()

        if not topic:
            topic = "General Overview Report"

        return {
            "doc_type": doc_type,
            "topic": topic,
            "raw": text
        }

    def generate_document_from_prompt(self, raw_text: str, llm_engine, model_name: str = "gemini-2.5-flash", user_name: str = "User") -> dict:
        """
        Orchestrates full document generation:
        1. Analyzes topic & format
        2. Prompts LLM for comprehensive, high-quality structured content
        3. Parses into title, subtitle, sections, bullets, and tables
        4. Compiles PDF or Word file
        5. Saves to Desktop and opens it
        """
        intent = self.detect_document_intent(raw_text)
        if not intent:
            return {"success": False, "error": "Could not identify document creation request."}

        doc_type = intent["doc_type"]
        topic = intent["topic"]

        # Prompt LLM for structured document outline
        author_name = f"ALIAS AI Assistant (for {user_name})"
        generation_prompt = f"""You are an elite executive document author and technical researcher.
Your task is to generate a comprehensive, highly thorough, publication-grade document about: "{topic}".

The document must contain:
1. A clear, professional Title and descriptive Subtitle.
2. An Executive Summary or Introduction.
3. 3 to 5 detailed, structured sections with comprehensive explanations, actionable insights, and bulleted takeaways.
4. At least one structured comparison or summary table with clear column headers (3-4 columns) and 3-5 rows.
5. A Conclusion and Strategic Recommendations section.

OUTPUT REQUIREMENTS:
Output ONLY valid JSON matching this schema:
{{
  "title": "Document Title",
  "subtitle": "Comprehensive Technical and Strategic Overview",
  "sections": [
    {{
      "heading": "1. Executive Summary",
      "paragraphs": [
        "In-depth paragraph 1...",
        "In-depth paragraph 2..."
      ],
      "bullets": [
        "Core Key Takeaway 1",
        "Core Key Takeaway 2",
        "Core Key Takeaway 3"
      ]
    }},
    {{
      "heading": "2. Comprehensive Analysis & Comparison",
      "paragraphs": [
        "Analysis paragraph detailing architectures or strategies..."
      ],
      "table": {{
        "headers": ["Category", "Core Specification", "Key Benefit", "Recommended Usage"],
        "rows": [
          ["Option A", "Spec details", "Primary benefit", "Best for enterprise"],
          ["Option B", "Spec details", "Secondary benefit", "Best for speed"],
          ["Option C", "Spec details", "Tertiary benefit", "Best for edge"]
        ]
      }},
      "bullets": []
    }},
    {{
      "heading": "3. Implementation Framework & Best Practices",
      "paragraphs": [
        "Detailed guidance on execution and practical methodology..."
      ],
      "bullets": [
        "Practice A: Key operational detail",
        "Practice B: Strategic monitoring"
      ]
    }},
    {{
      "heading": "4. Conclusion & Strategic Next Steps",
      "paragraphs": [
        "Final synthesis and concluding roadmap..."
      ],
      "bullets": [
        "Immediate Action Item 1",
        "Long-term Milestone 2"
      ]
    }}
  ]
}}
Do NOT output anything else before or after the JSON block.
"""

        doc_data = None
        try:
            llm_response = llm_engine.generate_response(generation_prompt, temperature=0.3, model_name=model_name)
            doc_data = self._parse_llm_json(llm_response)
        except Exception as e:
            print(f"[-] LLM Generation Error during document creation: {e}")

        # Fallback if LLM parsing failed
        if not doc_data:
            doc_data = self._create_fallback_doc_data(topic, author_name)
        else:
            doc_data["author"] = author_name

        # Sanitize filename
        safe_title = re.sub(r'[^a-zA-Z0-9_\- ]+', '', doc_data.get("title", topic)).strip().replace(" ", "_")
        if not safe_title or len(safe_title) < 3:
            safe_title = "ALIAS_Document"
        safe_title = safe_title[:40]

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_dir = self.get_output_directory()

        if doc_type == "pdf":
            filename = f"{safe_title}_{timestamp}.pdf"
            target_path = os.path.join(out_dir, filename)
            backup_path = str(self.local_docs_dir / filename)
            
            result = self.build_pdf(doc_data, target_path, backup_path)
        else:
            filename = f"{safe_title}_{timestamp}.docx"
            target_path = os.path.join(out_dir, filename)
            backup_path = str(self.local_docs_dir / filename)
            
            result = self.build_docx(doc_data, target_path, backup_path)

        # Attempt to launch the file automatically on Windows
        if result.get("success"):
            try:
                if hasattr(os, 'startfile'):
                    os.startfile(target_path)
            except Exception as oe:
                print(f"[-] Note: Could not auto-open file via startfile: {oe}")

        return result

    def _parse_llm_json(self, text: str) -> dict:
        """Parses structured JSON from LLM output, handling markdown fences."""
        if not text:
            return None
        cleaned = text.strip()
        if "```json" in cleaned:
            cleaned = cleaned.split("```json")[1].split("```")[0].strip()
        elif "```" in cleaned:
            cleaned = cleaned.split("```")[1].split("```")[0].strip()

        try:
            data = json.loads(cleaned)
            if isinstance(data, dict) and "sections" in data:
                return data
        except Exception as e:
            print(f"[-] JSON parse error: {e}")
        return None

    def _create_fallback_doc_data(self, topic: str, author_name: str) -> dict:
        """Provides high-quality structured content if LLM fails or is offline."""
        title = topic.title()
        return {
            "title": f"Report: {title}",
            "subtitle": "Prepared by ALIAS Intelligent Desktop Assistant",
            "author": author_name,
            "sections": [
                {
                    "heading": "1. Executive Summary",
                    "paragraphs": [
                        f"This document provides a comprehensive structured summary regarding {topic}.",
                        "It outlines the essential principles, key objectives, and operational recommendations."
                    ],
                    "bullets": [
                        "Identified high-priority functional objectives.",
                        "Comprehensive analysis of contemporary methodologies.",
                        "Clear roadmap for structured implementation."
                    ]
                },
                {
                    "heading": "2. Comparative Analysis",
                    "paragraphs": [
                        "Below is a structured overview of core parameters and assessment criteria."
                    ],
                    "table": {
                        "headers": ["Metric", "Baseline", "Target", "Status"],
                        "rows": [
                            ["Performance", "Nominal", "Optimized", "Active"],
                            ["Reliability", "98.5%", "99.9%", "Verified"],
                            ["Usability", "Standard", "Advanced", "Complete"]
                        ]
                    },
                    "bullets": []
                },
                {
                    "heading": "3. Strategic Recommendations",
                    "paragraphs": [
                        "To maximize value, adhere to documented best practices and maintain regular review cycles."
                    ],
                    "bullets": [
                        "Establish continuous monitoring metrics.",
                        "Align outcomes with long-term strategic targets."
                    ]
                }
            ]
        }

    def build_pdf(self, doc_data: dict, target_path: str, backup_path: str) -> dict:
        """Builds an elegant, publication-grade PDF using ReportLab."""
        try:
            doc = SimpleDocTemplate(
                target_path,
                pagesize=letter,
                leftMargin=54,
                rightMargin=54,
                topMargin=54,
                bottomMargin=54
            )

            styles = getSampleStyleSheet()

            # Custom typography and palette
            title_style = ParagraphStyle(
                'DocTitle',
                parent=styles['Title'],
                fontSize=22,
                leading=26,
                textColor=colors.HexColor('#0F172A'),
                fontName='Helvetica-Bold',
                alignment=0,
                spaceAfter=4
            )
            subtitle_style = ParagraphStyle(
                'DocSub',
                parent=styles['Normal'],
                fontSize=12,
                leading=16,
                textColor=colors.HexColor('#475569'),
                fontName='Helvetica-Oblique',
                spaceAfter=6
            )
            meta_style = ParagraphStyle(
                'DocMeta',
                parent=styles['Normal'],
                fontSize=9,
                leading=12,
                textColor=colors.HexColor('#64748B'),
                spaceAfter=12
            )
            h1_style = ParagraphStyle(
                'DocH1',
                parent=styles['Heading1'],
                fontSize=14,
                leading=18,
                textColor=colors.HexColor('#1E3A8A'),
                fontName='Helvetica-Bold',
                spaceBefore=14,
                spaceAfter=6,
                keepWithNext=True
            )
            body_style = ParagraphStyle(
                'DocBody',
                parent=styles['Normal'],
                fontSize=10,
                leading=14,
                textColor=colors.HexColor('#1E293B'),
                spaceAfter=6
            )
            bullet_style = ParagraphStyle(
                'DocBullet',
                parent=styles['Normal'],
                fontSize=10,
                leading=14,
                textColor=colors.HexColor('#334155'),
                leftIndent=18,
                firstLineIndent=-10,
                spaceAfter=4
            )

            story = [
                Paragraph(doc_data.get("title", "Document"), title_style),
                Paragraph(doc_data.get("subtitle", "Generated by ALIAS"), subtitle_style),
                Paragraph(f"Author: {doc_data.get('author', 'ALIAS')}  ·  Date: {datetime.now().strftime('%B %d, %Y')}", meta_style),
                HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#2563EB'), spaceAfter=14),
            ]

            sections = doc_data.get("sections", [])
            preview_snippets = []

            for sec in sections:
                heading = sec.get("heading", "")
                if heading:
                    story.append(Paragraph(heading, h1_style))
                
                for p in sec.get("paragraphs", []):
                    story.append(Paragraph(p, body_style))
                    if len(preview_snippets) < 2:
                        preview_snippets.append(p[:140] + "...")

                for b in sec.get("bullets", []):
                    story.append(Paragraph(f"&bull;  {b}", bullet_style))

                table_info = sec.get("table")
                if table_info and isinstance(table_info, dict) and "headers" in table_info:
                    headers = table_info.get("headers", [])
                    rows = table_info.get("rows", [])
                    if headers and rows:
                        th_style = ParagraphStyle('TH', parent=body_style, textColor=colors.white, fontName='Helvetica-Bold')
                        tbl_data = [[Paragraph(str(h), th_style) for h in headers]]
                        for row in rows:
                            tbl_data.append([Paragraph(str(cell), body_style) for cell in row])
                        
                        col_count = len(headers)
                        col_width = 504 / max(col_count, 1)

                        t = Table(tbl_data, colWidths=[col_width] * col_count)
                        t.setStyle(TableStyle([
                            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E3A8A')),
                            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
                            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                            ('BOTTOMPADDING', (0,0), (-1,-1), 5),
                            ('TOPPADDING', (0,0), (-1,-1), 5),
                            ('LEFTPADDING', (0,0), (-1,-1), 6),
                            ('RIGHTPADDING', (0,0), (-1,-1), 6),
                            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
                            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')])
                        ]))
                        story.append(Spacer(1, 6))
                        story.append(t)
                        story.append(Spacer(1, 8))

            doc.build(story)

            # Copy to backup uploads
            try:
                import shutil
                shutil.copy2(target_path, backup_path)
            except Exception:
                pass

            file_size_kb = os.path.getsize(target_path) / 1024

            return {
                "success": True,
                "type": "PDF",
                "title": doc_data.get("title"),
                "filename": os.path.basename(target_path),
                "file_path": target_path,
                "backup_path": backup_path,
                "file_size_kb": file_size_kb,
                "sections_count": len(sections),
                "preview": "\n".join(preview_snippets)
            }
        except Exception as e:
            print(f"[-] PDF Build Error: {e}")
            import traceback
            traceback.print_exc()
            return {"success": False, "error": f"Failed to build PDF: {str(e)}"}

    def build_docx(self, doc_data: dict, target_path: str, backup_path: str) -> dict:
        """Builds a beautifully formatted Microsoft Word (.docx) document."""
        try:
            doc = docx.Document()

            # Set standard margins (1 inch)
            for section in doc.sections:
                section.top_margin = Inches(1)
                section.bottom_margin = Inches(1)
                section.left_margin = Inches(1)
                section.right_margin = Inches(1)

            # Title
            title_p = doc.add_paragraph()
            t_run = title_p.add_run(doc_data.get("title", "Document"))
            t_run.font.name = 'Segoe UI'
            t_run.font.size = Pt(22)
            t_run.font.bold = True
            t_run.font.color.rgb = RGBColor(15, 23, 42)
            title_p.paragraph_format.space_after = Pt(2)

            # Subtitle
            sub_p = doc.add_paragraph()
            s_run = sub_p.add_run(doc_data.get("subtitle", "Prepared by ALIAS"))
            s_run.font.name = 'Segoe UI'
            s_run.font.size = Pt(12)
            s_run.font.italic = True
            s_run.font.color.rgb = RGBColor(71, 85, 105)
            sub_p.paragraph_format.space_after = Pt(4)

            # Metadata
            meta_p = doc.add_paragraph()
            m_run = meta_p.add_run(f"Author: {doc_data.get('author', 'ALIAS')}  |  Date: {datetime.now().strftime('%B %d, %Y')}")
            m_run.font.name = 'Segoe UI'
            m_run.font.size = Pt(9)
            m_run.font.color.rgb = RGBColor(100, 116, 139)
            meta_p.paragraph_format.space_after = Pt(14)

            sections = doc_data.get("sections", [])
            preview_snippets = []

            for sec in sections:
                heading = sec.get("heading", "")
                if heading:
                    h_p = doc.add_heading(level=1)
                    h_run = h_p.add_run(heading)
                    h_run.font.name = 'Segoe UI'
                    h_run.font.size = Pt(14)
                    h_run.font.bold = True
                    h_run.font.color.rgb = RGBColor(30, 58, 138)
                    h_p.paragraph_format.space_before = Pt(14)
                    h_p.paragraph_format.space_after = Pt(4)

                for p in sec.get("paragraphs", []):
                    par = doc.add_paragraph()
                    p_run = par.add_run(p)
                    p_run.font.name = 'Segoe UI'
                    p_run.font.size = Pt(10.5)
                    p_run.font.color.rgb = RGBColor(30, 41, 59)
                    par.paragraph_format.space_after = Pt(6)
                    par.paragraph_format.line_spacing = 1.15
                    if len(preview_snippets) < 2:
                        preview_snippets.append(p[:140] + "...")

                for b in sec.get("bullets", []):
                    bp = doc.add_paragraph(style='List Bullet')
                    b_run = bp.add_run(b)
                    b_run.font.name = 'Segoe UI'
                    b_run.font.size = Pt(10)
                    b_run.font.color.rgb = RGBColor(51, 65, 85)
                    bp.paragraph_format.space_after = Pt(3)

                table_info = sec.get("table")
                if table_info and isinstance(table_info, dict) and "headers" in table_info:
                    headers = table_info.get("headers", [])
                    rows = table_info.get("rows", [])
                    if headers and rows:
                        tbl = doc.add_table(rows=1, cols=len(headers))
                        tbl.style = 'Table Grid'
                        hdr_cells = tbl.rows[0].cells
                        for i, h_text in enumerate(headers):
                            hdr_cells[i].text = str(h_text)
                            for run in hdr_cells[i].paragraphs[0].runs:
                                run.font.bold = True
                                run.font.name = 'Segoe UI'
                                run.font.color.rgb = RGBColor(30, 58, 138)

                        for row_data in rows:
                            row_cells = tbl.add_row().cells
                            for i, val in enumerate(row_data):
                                row_cells[i].text = str(val)
                                for run in row_cells[i].paragraphs[0].runs:
                                    run.font.name = 'Segoe UI'
                                    run.font.size = Pt(9.5)
                        
                        doc.add_paragraph().paragraph_format.space_after = Pt(8)

            doc.save(target_path)

            try:
                import shutil
                shutil.copy2(target_path, backup_path)
            except Exception:
                pass

            file_size_kb = os.path.getsize(target_path) / 1024

            return {
                "success": True,
                "type": "Word (DOCX)",
                "title": doc_data.get("title"),
                "filename": os.path.basename(target_path),
                "file_path": target_path,
                "backup_path": backup_path,
                "file_size_kb": file_size_kb,
                "sections_count": len(sections),
                "preview": "\n".join(preview_snippets)
            }
        except Exception as e:
            print(f"[-] DOCX Build Error: {e}")
            import traceback
            traceback.print_exc()
            return {"success": False, "error": f"Failed to build DOCX: {str(e)}"}
